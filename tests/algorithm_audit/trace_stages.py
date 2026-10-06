"""Capture production stage results without changing their return values."""
from pathlib import Path
import sys
import pandas as pd
from bom_categorizer import main as pipeline
from docx import Document
OUT = Path('docs/algorithm_audit_evidence')
aggregate = pipeline.aggregate_duplicate_items
classify = pipeline.run_classification

def capture_aggregate(*args, **kwargs):
    args[0].to_csv(OUT/'MKVH_before_aggregate.csv',index=False,encoding='utf-8-sig')
    result = aggregate(*args, **kwargs)
    result.to_csv(OUT/'MKVH_after_aggregate.csv',index=False,encoding='utf-8-sig')
    return result

def capture_classify(*args, **kwargs):
    result = classify(*args, **kwargs)
    result.to_csv(OUT/'MKVH_after_classification.csv',index=False,encoding='utf-8-sig')
    print('CLASSIFICATION COUNTS',result.category.value_counts().to_dict())
    print('NON_BOM',result.loc[result.category.eq('non_bom')].to_dict('records'))
    return result

pipeline.aggregate_duplicate_items = capture_aggregate
pipeline.run_classification = capture_classify
sys.argv = ['audit','--inputs',str(OUT/'plata_MKVH-02_audit.docx'),'--xlsx',str(OUT/'plata_MKVH-02_audit_bom.xlsx'),'--combine','--no-interactive','--assign-json',str(OUT/'no_rules.json')]
pipeline.main()
doc = Document(OUT/'plata_MKVH-02_audit.docx')
import json
rows = [{'table':ti,'row':ri,'cells':[c.text for c in row.cells]} for ti,t in enumerate(doc.tables,1) for ri,row in enumerate(t.rows,1)]
(OUT/'plata_MKVH-02_raw.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding='utf-8')
print('SOURCE ROWS',len(rows),'TABLES',len(doc.tables))
