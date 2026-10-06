"""Regression checks across export, allocation, and repeat RKM processing."""
import pandas as pd
import pytest

from bom_categorizer.excel_writer import format_excel_output
from bom_categorizer.tru_merger import (
    prepare_merge_context, merge_tru_into_bom, build_ostatki_and_zapas_reports,
    find_matching_tru_row,
)
from bom_categorizer.tru_rkm_processor import process_rkm_file
from bom_categorizer.utils import parse_number


def test_final_export_preserves_distinct_mpns():
    source = pd.DataFrame({'description': ['Резистор', 'Резистор'],
                           'partnumber': ['ABC123-A', 'ABC123-B'],
                           'qty': [2, 3], 'reference': ['R1', 'R2']})
    output = format_excel_output(source, 'Резисторы', 'description')
    assert len(output) == 2
    assert set(output['partnumber']) == {'ABC123-A', 'ABC123-B'}
    assert output['шт.'].sum() == 5


def supply(quantity=5, cost=50):
    return pd.DataFrame({'Наименование': ['ABC123'], 'Артикул': ['101'],
                         'Количество': [quantity], 'Стоимость': [cost]})


def test_allocation_conserves_supply_across_sheets():
    frames = [pd.DataFrame({'Наименование ИВП': ['ABC123'], 'шт.': [q]}) for q in [2, 3]]
    trus = [supply()]
    context = prepare_merge_context(frames, trus)
    costs = []
    for frame in frames:
        result, matched, used = merge_tru_into_bom(frame, trus, allocation_context=context)
        shortage, surplus = build_ostatki_and_zapas_reports(result, matched)
        assert shortage.empty and surplus.empty and used == {0}
        costs.append(result.iloc[0]['Стоимость'])
    assert costs == [20, 30]


def test_invalid_consumer_does_not_hide_surplus():
    frame = pd.DataFrame({'Наименование ИВП': ['ABC123', 'ABC123'], 'шт.': [2, '?']})
    result, matched, _ = merge_tru_into_bom(frame, [supply()])
    _, surplus = build_ostatki_and_zapas_reports(result, matched)
    assert matched == {0}
    assert surplus['шт.'].sum() == 3
    assert 'проверка количества' in result.iloc[1]['Статус сопоставления']


def test_rounding_does_not_create_or_destroy_cents():
    frame = pd.DataFrame({'Наименование ИВП': ['ABC123'] * 3, 'шт.': [1] * 3})
    result, _, _ = merge_tru_into_bom(frame, [supply(3, 1)])
    assert sorted(result['Стоимость'].tolist()) == [0.33, 0.33, 0.34]
    assert result['Стоимость'].sum() == pytest.approx(1)


def test_conflicting_supply_units_require_review():
    trus = pd.concat([supply(), supply()], ignore_index=True)
    trus['Единица измерения'] = ['шт.', 'м']
    assert find_matching_tru_row('ABC123', '', trus) is None


def test_repeat_rkm_preserves_first_item_and_unknown_units(tmp_path):
    source, output = tmp_path / 'rkm.xlsx', tmp_path / 'result.xlsx'
    pd.DataFrame({'invalid index heading': ['text', 2],
                  'Наименование': ['ABC123', 'DEF456'], 'шт.': [2, 1.5],
                  'Цена': [10, 20], 'Стоимость': [20, 30],
                  'Документы': ['Документ 1', 'Документ 2'],
                  'Поставщик': ['А', 'Б']}).to_excel(source, index=False)
    ok, message = process_rkm_file(str(source), str(output))
    assert ok, message
    result = pd.read_excel(output).dropna(subset=['Наименование'])
    assert set(result['Наименование']) == {'ABC123', 'DEF456'}
    assert result['Количество'].sum() == 3.5
    assert result['Стоимость'].sum() == 50
    assert result['Единица измерения'].isna().all()
    assert 'нет единиц измерения' in message


def test_rkm_explicit_zero_does_not_use_historical_quantity(tmp_path):
    row = [None] * 24
    row[0:2] = ['2.1.1', 'ABC123']
    row[5], row[6], row[8], row[10] = 'шт.', 10, 12, 120
    row[17:20] = [0, 0, 0]
    source, output = tmp_path / 'rkm.xlsx', tmp_path / 'result.xlsx'
    pd.DataFrame([['№ п/п', 'Наименование'] + [None]*22, row]).to_excel(source, header=False, index=False)
    ok, message = process_rkm_file(str(source), str(output))
    assert ok, message
    result = pd.read_excel(output)
    assert result.iloc[0]['Количество'] == 0
    assert result.iloc[0]['Стоимость'] == 0


def test_numeric_scientific_representation_is_not_rejected():
    assert parse_number(1e-7) == 1e-7
    assert parse_number(float('inf')) is None


def test_missing_word_input_does_not_produce_partial_result(tmp_path):
    from bom_categorizer.main import load_and_combine_inputs
    with pytest.raises(ValueError, match='Word'):
        load_and_combine_inputs([str(tmp_path / 'missing.docx')])


@pytest.mark.parametrize('title', ['Лист регистрации изменений', 'Изм. Номера листов (страниц)'])
def test_word_change_register_is_not_a_component_table(tmp_path, title):
    from docx import Document
    from bom_categorizer.parsers import parse_docx
    document = Document()
    parts = document.add_table(rows=2, cols=3)
    for cell, value in zip(parts.rows[0].cells, ['Поз.', 'Наименование', 'Количество']):
        cell.text = value
    for cell, value in zip(parts.rows[1].cells, ['R1', 'ABC123', '2']):
        cell.text = value
    changes = document.add_table(rows=3, cols=4)
    changes.cell(0, 0).text = title
    for cell, value in zip(changes.rows[1].cells, ['Изм.', 'Лист', 'Подп.', 'Дата']):
        cell.text = value
    path = tmp_path / 'bom.docx'
    document.save(path)
    result = parse_docx(str(path))
    assert len(result) == 1 and result.iloc[0]['qty'] == 2


def test_tru_unit_survives_processing_and_mismatch_is_rejected(tmp_path):
    from bom_categorizer.tru_rkm_processor import process_tru_files_batch, _read_tru_file
    source, output = tmp_path / 'tru.xlsx', tmp_path / 'processed.xlsx'
    trus = supply()
    trus['Единица измерения'], trus['Цена'] = 'м', 10
    trus.to_excel(source, index=False)
    assert process_tru_files_batch([str(source)], str(output))[0]
    reread = _read_tru_file(str(output))
    assert reread.iloc[0]['Единица измерения'] == 'м'
    bom = pd.DataFrame({'Наименование ИВП': ['ABC123'], 'шт.': [1], 'unit': ['шт.']})
    result, matched, used = merge_tru_into_bom(bom, [reread])
    assert not matched and not used
    assert 'единиц' in result.iloc[0]['Статус сопоставления']


def test_rkm_keeps_different_prices_separate(tmp_path):
    source, output = tmp_path / 'rkm.xlsx', tmp_path / 'result.xlsx'
    pd.DataFrame({'Наименование': ['ABC123', 'ABC123'], 'Количество': [1, 2],
                  'Цена': [10, 20], 'Стоимость': [10, 40], 'Поставщик': ['A', 'A'],
                  'Единица измерения': ['шт.', 'шт.']}).to_excel(source, index=False)
    assert process_rkm_file(str(source), str(output))[0]
    result = pd.read_excel(output).dropna(subset=['Наименование'])
    assert len(result) == 2
    assert result['Стоимость'].sum() == 50


@pytest.mark.parametrize('value', [
    'Фильтр де1', 'Фильтр ДЕ-2', 'Фильтр ДЕ 5.067.066-03',
    'Фильтр DE3', 'Фильтр е5.067.066-03', 'Фильтр ДЕ пять',
])
def test_all_development_code_spellings_are_our_developments(value):
    from bom_categorizer.classifiers import classify_row
    from bom_categorizer.formatters import extract_tu_code
    assert classify_row('C1', value, None, None, strict=True) == 'our_developments'
    name, code = extract_tu_code(value)
    assert code and ('фильтр' in name.lower() or name.lower() == value.lower())


def test_development_code_is_kept_in_output_column():
    from bom_categorizer.excel_writer import format_excel_output
    source = pd.DataFrame({'description': ['Фильтр ДЕ-2'], 'qty': [1], 'reference': ['C1']})
    output = format_excel_output(source, 'Наши разработки', 'description')
    assert len(output) == 1
    assert 'ДЕ' in str(output.iloc[0]['ТУ/Производитель']).upper()


def test_development_code_overrides_stale_input_category():
    from bom_categorizer.main import run_classification
    source = pd.DataFrame({'description': ['Фильтр ДЕ-2'], 'category': ['capacitors']})
    result = run_classification(source, None, 'description', None, None, False)
    assert result.iloc[0]['category'] == 'our_developments'


@pytest.mark.parametrize('value', [
    'Фильтр ¶e 5.067.066-03',
    'Фильтр 𝑑e 5.067.066-03',
    'Фильтр e5.067.066-03',
])
def test_word_symbol_font_de_prefix_is_restored(value):
    from bom_categorizer.formatters import clean_component_name, extract_tu_code
    from bom_categorizer.classifiers import classify_row
    cleaned = clean_component_name(value)
    assert classify_row('C1', cleaned, None, None, strict=True) == 'our_developments'
    name, code = extract_tu_code(cleaned)
    assert code.lower().replace('е', 'e').startswith(('де', 'de', 'дe', 'de')) and '5.067.066' in code.lower()
