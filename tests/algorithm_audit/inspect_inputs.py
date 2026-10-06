"""Read-only inspection of user-supplied inputs. Outputs stay in evidence dir."""
import hashlib
import json
from pathlib import Path
import pandas as pd
from docx import Document
from bom_categorizer.parsers import parse_docx
from bom_categorizer.tru_rkm_processor import _read_tru_file, process_rkm_file

OUT = Path('docs/algorithm_audit_evidence')
BU = Path(r'C:\Project\!SHSK-M\BU')
TPY = Path(r'C:\Project\!SHSK-M\TPY\ТРУ_17043_2025')
paths = [BU / n for n in ['Nashi4BU.docx', 'Plata_control.docx', 'plata_MKVH-02.doc']]
paths += [TPY / n for n in ['PKM_2026.xlsx', 'ТРУ.953033.17043.xls', 'ТРУ.953033.17043_002.xls']]
manifest = []
for path in paths:
    binary = path.read_bytes()
    manifest.append({'path': str(path), 'sha256': hashlib.sha256(binary).hexdigest(), 'size': len(binary)})
    print('\nFILE', path.name, 'magic', binary[:12].hex())
    if path.suffix == '.docx':
        doc = Document(path)
        raw = []
        for ti, table in enumerate(doc.tables, 1):
            for ri, row in enumerate(table.rows, 1):
                raw.append({'table': ti, 'row': ri, 'cells': [c.text for c in row.cells]})
        (OUT / (path.stem + '_raw.json')).write_text(json.dumps(raw, ensure_ascii=False, indent=2), encoding='utf-8')
        df = parse_docx(str(path))
        df.to_csv(OUT / (path.stem + '_parsed.csv'), index=False, encoding='utf-8-sig')
        print('tables', len(doc.tables), 'raw_rows', len(raw), 'parsed', len(df), 'qty', df.get('qty', pd.Series(dtype=float)).sum())
        print(df.head(4).to_string(index=False))
    elif path.suffix in ('.xls', '.xlsx'):
        try:
            book = pd.ExcelFile(path)
            print('sheets', book.sheet_names)
            for sheet in book.sheet_names:
                df = pd.read_excel(book, sheet_name=sheet, header=None)
                print('sheet', sheet, 'shape', df.shape)
                print(df.head(10).to_string(index=False, header=False))
                df.to_csv(OUT / (path.stem + '_' + str(book.sheet_names.index(sheet)) + '_raw.csv'), index=False, header=False, encoding='utf-8-sig')
            if path.suffix == '.xls':
                parsed = _read_tru_file(str(path))
                print('parsed TRU', None if parsed is None else len(parsed))
                if parsed is not None:
                    parsed.to_csv(OUT / (path.stem + '_parsed.csv'), index=False, encoding='utf-8-sig')
                    print(parsed.head(3).to_string(index=False))
            else:
                print('RKM process', process_rkm_file(str(path), str(OUT / 'PKM_audit_result.xlsx')))
        except Exception as exc:
            print(type(exc).__name__, str(exc))
    else:
        try:
            parse_docx(str(path))
        except Exception as exc:
            print('direct DOC parser:', type(exc).__name__, str(exc))
(OUT / 'input_manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
